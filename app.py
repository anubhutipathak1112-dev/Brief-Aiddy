import streamlit as st
import os
import json
import tempfile
import textwrap
import google.generativeai as genai
from pydantic import BaseModel
from typing import List
import PyPDF2
import docx
from fpdf import FPDF

# --- 1. CONFIGURATION ---
os.environ["GEMINI_API_KEY"] = "AIzaSyAD7vinOWSWnDTvYkV0oCfUad9PNO2JAmk"
genai.configure(api_key=os.environ["GEMINI_API_KEY"])

# --- 2. ENTERPRISE DATA SCHEMA ---
class BriefAnalysis(BaseModel):
    brief_quality_score: int
    executive_summary: str
    action_plan: List[str]
    identified_gaps: List[str]
    required_agents: List[str] 
    mnc_project_timeline: List[str] 
    compliance_brand_assessment: str 
    client_followup_questions: List[str]

# --- 3. CORE AI LOGIC ---
def extract_text_from_file(file_path: str) -> str:
    ext = file_path.lower().split('.')[-1]
    text = ""
    if ext == 'txt':
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
    elif ext == 'pdf':
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            text = " ".join([page.extract_text() for page in reader.pages if page.extract_text() or ""])
    elif ext == 'docx':
        doc = docx.Document(file_path)
        text = " ".join([para.text for para in doc.paragraphs])
    return text

def process_brief_with_aiddy(raw_text: str):
    model = genai.GenerativeModel(
        model_name="gemini-2.5-flash",
        system_instruction="""You are 'Brief Aiddy', the Lead Intake & Orchestration AI for an enterprise digital marketing platform. 
        Analyze the client brief and output a strict JSON structure.
        
        1. Score quality (0-100) and write an executive summary.
        2. Detail a plan of action and identify critical missing information (gaps).
        3. List required downstream agents ONLY from this exact list: [GEOAiddy, MediaAiddy, Paid Media Aiddy, Social listening Aiddy, Content Aiddy, Design Aiddy].
        4. Create an 'mnc_project_timeline': Structure this for a large enterprise (e.g., Include phases like Stakeholder Alignment, Legal/Compliance Review, Sprint Execution, QA, and Deployment).
        5. Provide a 'compliance_brand_assessment': Briefly flag any potential regulatory, brand safety, or legal risks based on the brief's industry or requests.
        6. Generate polite client follow-up questions to resolve the gaps."""
    )

    prompt = f"Analyze the following client brief:\n\n{raw_text}"

    response = model.generate_content(
        prompt,
        generation_config=genai.GenerationConfig(
            response_mime_type="application/json",
            response_schema=BriefAnalysis,
            temperature=0.2, 
        ),
    )
    return json.loads(response.text)

# --- 4. BULLETPROOF PDF GENERATION LOGIC ---
def generate_pdf_report(result, email_body):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    def clean_text(text):
        # Replace smart quotes that often cause encoding crashes
        text = str(text).replace('”', '"').replace('“', '"').replace('’', "'").replace('–', '-')
        return text.encode('latin-1', 'replace').decode('latin-1')

    # Custom writer that manually slices lines to prevent fpdf crashes
    def safe_write(text):
        text_str = clean_text(text)
        if not text_str.strip():
            pdf.ln(7)
            return
        # Force a hard break at 85 chars, no matter what
        lines = textwrap.wrap(text_str, width=85, break_long_words=True)
        for line in lines:
            pdf.cell(0, 7, line, ln=True)

    pdf.set_font("Helvetica", style="B", size=16)
    pdf.cell(0, 10, "Brief Aiddy - Enterprise Intelligence Report", ln=True, align="C")
    pdf.ln(5)

    sections = [
        ("Brief Readiness Score", f"{result.get('brief_quality_score', 0)}/100"),
        ("Executive Summary", result.get("executive_summary", "N/A")),
        ("Action Plan", result.get("action_plan", [])),
        ("Identified Gaps", result.get("identified_gaps", [])),
        ("Required Specialized Agents", result.get("required_agents", [])),
        ("Enterprise Project Timeline", result.get("mnc_project_timeline", [])),
        ("Compliance & Brand Assessment", result.get("compliance_brand_assessment", "N/A")),
        ("Drafted Client Communication", email_body)
    ]

    for title, content in sections:
        pdf.set_font("Helvetica", style="B", size=12)
        pdf.cell(0, 10, clean_text(title), ln=True)
        pdf.set_font("Helvetica", size=11)
        
        if isinstance(content, list):
            for item in content:
                safe_write(f"- {item}")
        else:
            # Handle AI outputs that include line breaks natively
            for para in str(content).split('\n'):
                safe_write(para)
        pdf.ln(5)

    return bytes(pdf.output())

# --- 5. VIBRANT PROFESSIONAL UI & STYLING ---
st.set_page_config(page_title="Brief Aiddy | Orchestration", layout="wide")

st.markdown("""
    <style>
    .block-container { padding-top: 2rem; padding-bottom: 2rem; }
    .main { background-color: #FAFAFA; }
    
    .stTextArea textarea { border-radius: 8px; border: 1.5px solid #E2E8F0; font-size: 1.05rem; padding: 12px;}
    .stTextArea textarea:focus { border-color: #3B82F6; box-shadow: 0 0 0 2px rgba(59, 130, 246, 0.2); }
    
    .stTabs [data-baseweb="tab-list"] { 
        gap: 12px; 
        background-color: #F1F5F9; 
        padding: 8px; 
        border-radius: 12px;
        margin-bottom: 20px;
    }
    .stTabs [data-baseweb="tab"] { 
        background-color: #FFFFFF; 
        border: 1px solid #E2E8F0; 
        border-radius: 8px; 
        padding: 14px 20px; 
        font-weight: 600; 
        font-size: 1.1rem; 
        color: #475569;
        transition: all 0.2s ease-in-out;
    }
    .stTabs [data-baseweb="tab"]:hover {
        border-color: #3B82F6;
        color: #3B82F6;
    }
    .stTabs [aria-selected="true"] { 
        background-color: #3B82F6 !important; 
        color: #FFFFFF !important; 
        border: none;
        box-shadow: 0 4px 6px -1px rgba(59, 130, 246, 0.4);
    }
    
    .feature-card {
        border: 1px solid #E2E8F0;
        border-radius: 10px;
        padding: 20px;
        background-color: #FFFFFF;
        text-align: center;
        font-weight: 700;
        color: #1E293B;
        font-size: 1.15rem;
        box-shadow: 0 2px 4px rgba(0,0,0,0.04);
        margin-bottom: 16px;
        transition: transform 0.2s ease;
    }
    .feature-card:hover { transform: translateY(-2px); box-shadow: 0 4px 8px rgba(0,0,0,0.08); }
    </style>
""", unsafe_allow_html=True)

st.markdown("<h1 style='font-size: 3.5rem; color: #0F172A; margin-bottom: 0px; padding-bottom: 0px;'>Brief Aiddy <span style='color: #3B82F6;'>.</span></h1>", unsafe_allow_html=True)
st.markdown("<p style='font-size: 1.25rem; color: #64748B; margin-top: 5px; font-weight: 500;'>Intelligent Intake, Compliance Guardian & Agent Orchestration</p>", unsafe_allow_html=True)
st.markdown("<hr style='margin-top: 10px; margin-bottom: 30px; border-color: #E2E8F0;'>", unsafe_allow_html=True)

col_input, col_output = st.columns([1, 1.4], gap="large")

with col_input:
    st.markdown("<h3 style='color: #0F172A; margin-bottom: 15px;'>📝 Input Content</h3>", unsafe_allow_html=True)
    input_type = st.radio("Content Source", ["Paste Text", "Upload Document"], horizontal=True, label_visibility="collapsed")
    
    raw_brief_text = ""
    
    if input_type == "Paste Text":
        raw_brief_text = st.text_area("Paste your unoptimized brief here...", height=350, 
                                      placeholder="Example: We need a Q3 campaign targeting millennials for our new fintech app. Budget is 150k...")
    else:
        st.info("Supported formats: PDF, DOCX, TXT")
        uploaded_file = st.file_uploader("Upload Brief Document", type=['txt', 'pdf', 'docx'], label_visibility="collapsed")
        if uploaded_file is not None:
            with tempfile.NamedTemporaryFile(delete=False, suffix=f".{uploaded_file.name.split('.')[-1]}") as tmp_file:
                tmp_file.write(uploaded_file.getvalue())
                tmp_path = tmp_file.name
            try:
                raw_brief_text = extract_text_from_file(tmp_path)
                st.success(f"Successfully extracted {len(raw_brief_text.split())} words.")
            except Exception as e:
                st.error("Error reading file.")
            finally:
                os.remove(tmp_path)

    st.write("")
    analyze_button = st.button("Optimize with Brief Aiddy", type="primary", use_container_width=True)

with col_output:
    st.markdown("<h3 style='color: #0F172A; margin-bottom: 15px;'><\> Optimized Output</h3>", unsafe_allow_html=True)
    
    if not analyze_button:
        st.markdown("<div style='padding-top: 30px; padding-bottom: 30px;'></div>", unsafe_allow_html=True)
        st.markdown("<h3 style='text-align: center; color: #334155; margin-bottom: 10px;'>Ready to Optimize</h3>", unsafe_allow_html=True)
        st.markdown("<p style='text-align: center; color: #64748B; font-size: 1.15rem; margin-bottom: 40px;'>Paste your content on the left and click <b style='color:#3B82F6;'>Optimize with Brief Aiddy</b>.</p>", unsafe_allow_html=True)
        
        cols = st.columns(2)
        with cols[0]:
            st.markdown("<div class='feature-card'>🎯 Action Plan & Gaps</div>", unsafe_allow_html=True)
            st.markdown("<div class='feature-card'>🤖 Agent Routing</div>", unsafe_allow_html=True)
        with cols[1]:
            st.markdown("<div class='feature-card'>⏱️ MNC Project Timeline</div>", unsafe_allow_html=True)
            st.markdown("<div class='feature-card'>🛡️ Compliance Guardian</div>", unsafe_allow_html=True)

    if analyze_button and raw_brief_text:
        with st.spinner("Analyzing compliance, forecasting timelines, and routing agents..."):
            try:
                result = process_brief_with_aiddy(raw_brief_text)
                
                email_body = "Hi Team,\n\nThanks for the brief. To ensure compliance and accurate timeline forecasting, could you clarify:\n\n"
                for q in result.get("client_followup_questions", []):
                    email_body += f"- {q}\n"
                email_body += "\nBest,\nBrief Aiddy"
                
                tab1, tab2, tab3, tab4, tab5 = st.tabs([
                    "📊 Exec Summary", 
                    "🎯 Plan & Gaps", 
                    "🤖 Aiddy Agents", 
                    "⏱️ Timeline", 
                    "🛡️ Compliance"
                ])
                
                with tab1:
                    score = result.get('brief_quality_score', 0)
                    c1, c2 = st.columns([1, 2])
                    c1.metric("Brief Readiness Score", f"{score}/100")
                    c2.markdown("#### Executive Summary")
                    c2.write(result.get("executive_summary", "N/A"))

                with tab2:
                    st.markdown("#### Plan of Actions")
                    for action in result.get("action_plan", []):
                        st.write(f"✅ {action}")
                    
                    st.divider()
                    st.markdown("#### Identified Gaps (Requires Client Comms)")
                    for gap in result.get("identified_gaps", []):
                        st.error(f"⚠️ {gap}")

                with tab3:
                    st.markdown("#### Required Specialized Agents")
                    agents = result.get("required_agents", [])
                    if agents:
                        acols = st.columns(2)
                        for i, agent in enumerate(agents):
                            acols[i % 2].success(f"**{agent}** assigned")
                    else:
                        st.write("No specialized agents identified.")

                with tab4:
                    st.markdown("#### Enterprise Project Timeline")
                    for step in result.get("mnc_project_timeline", []):
                        st.info(f"📅 {step}")

                with tab5:
                    st.markdown("#### Compliance & Brand Voice Guardian")
                    st.warning(result.get("compliance_brand_assessment", "No compliance data generated."))
                    
                    st.divider()
                    st.markdown("#### Auto-Generated Client Comms")
                    st.text_area("Review & Edit", value=email_body, height=200, label_visibility="collapsed")

                st.markdown("<hr style='margin-top: 30px; margin-bottom: 20px;'>", unsafe_allow_html=True)
                
                # The stabilized PDF generator
                pdf_bytes = generate_pdf_report(result, email_body)
                st.download_button(
                    label="📥 Download Enterprise Report (PDF)",
                    data=pdf_bytes,
                    file_name="Brief_Aiddy_Intelligence_Report.pdf",
                    mime="application/pdf",
                    type="primary",
                    use_container_width=True
                )

            except Exception as e:
                st.error(f"An error occurred during processing: {e}")
    elif analyze_button and not raw_brief_text:
        st.warning("Please paste text or upload a document first.")